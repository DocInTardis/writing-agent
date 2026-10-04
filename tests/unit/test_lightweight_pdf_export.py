from __future__ import annotations

from types import SimpleNamespace
import io

from docx import Document

from writing_agent.web.services.export_service import ExportService


def test_lightweight_pdf_fallback_writes_unicode_pdf(tmp_path) -> None:
    session = SimpleNamespace(
        document_v3={
            "sections": [
                {
                    "headerFooter": {
                        "header": [{"content": [{"type": "text", "text": "论文页眉"}]}],
                        "footer": [{"content": [{"type": "text", "text": "论文页脚"}]}],
                    }
                }
            ]
        }
    )
    output = tmp_path / "fallback.pdf"
    text = "毕业论文标题\n\n" + "中文正文用于验证无 Office 环境的 PDF 导出。\n\n" * 180

    ExportService._render_lightweight_pdf(text, output, session)

    payload = output.read_bytes()
    assert payload.startswith(b"%PDF")
    assert len(payload) > 10_000
    assert payload.rstrip().endswith(b"%%EOF")


def test_lightweight_pdf_uses_document_v3_page_geometry_and_numbering() -> None:
    session = SimpleNamespace(
        document_v3={
            "sections": [
                {
                    "layout": {
                        "pageSize": "A5",
                        "orientation": "landscape",
                        "marginTopMm": 20,
                        "marginRightMm": 15,
                        "marginBottomMm": 18,
                        "marginLeftMm": 25,
                    },
                    "headerFooter": {
                        "pageNumber": {
                            "enabled": True,
                            "startAt": 3,
                            "format": "upperRoman",
                            "position": "header",
                            "alignment": "right",
                        }
                    },
                }
            ]
        }
    )

    layout = ExportService._lightweight_pdf_layout(session)

    assert round(float(layout["width"]), 2) == 595.28
    assert round(float(layout["height"]), 2) == 419.53
    assert round(float(layout["margin_left"]), 2) == 70.87
    assert layout["page_number_start"] == 3
    assert layout["page_number_format"] == "upperRoman"
    assert layout["page_number_position"] == "header"
    assert layout["page_number_alignment"] == "right"


def test_document_v3_custom_style_is_projected_into_word_docx() -> None:
    document = Document()
    document.add_paragraph("需要按样式导出的正文")
    # The legacy parsed exporter can repeat the first V3 block as a synthetic
    # title.  The V3 projection pass must collapse that exporter artifact.
    document.add_paragraph("需要按样式导出的正文")
    source = io.BytesIO()
    document.save(source)
    session = SimpleNamespace(
        document_v3={
            "styles": [
                {
                    "id": "normal",
                    "name": "正文",
                    "kind": "paragraph",
                    "visible": True,
                    "properties": {"fontFamily": "SimSun", "fontSizePt": 12, "bold": True},
                },
                {
                    "id": "thesis-body",
                    "name": "论文正文",
                    "kind": "paragraph",
                    "basedOn": "normal",
                    "nextStyle": "thesis-body",
                    "visible": True,
                    "properties": {
                        "fontFamily": "Microsoft YaHei",
                        "fontSizePt": 11,
                        "bold": False,
                        "letterSpacingPt": 0.5,
                        "lineSpacing": 1.5,
                        "firstLineIndentEm": 2,
                        "borderStyle": "solid",
                        "borderColor": "#123456",
                        "borderWidthPt": 1,
                    },
                },
            ],
            "sections": [
                {
                    "content": [
                        {
                            "id": "p1",
                            "type": "paragraph",
                            "styleId": "thesis-body",
                            "content": [{"type": "text", "text": "需要按样式导出的正文"}],
                        }
                    ]
                }
            ],
        }
    )

    result = ExportService._apply_document_v3_styles(source.getvalue(), session)
    exported = Document(io.BytesIO(result))

    assert [paragraph.text for paragraph in exported.paragraphs] == ["需要按样式导出的正文"]
    assert exported.paragraphs[0].style.name == "论文正文"
    assert exported.styles["论文正文"].font.name == "Microsoft YaHei"
    assert exported.styles["论文正文"].font.size.pt == 11
    assert exported.styles["论文正文"].font.bold is False
    assert exported.styles["论文正文"].paragraph_format.line_spacing == 1.5
    style_xml = exported.styles["论文正文"]._element.xml
    assert 'w:spacing w:val="10"' in style_xml
    assert 'w:color="123456"' in style_xml


def test_document_v3_multilevel_numbering_is_projected_into_word_docx() -> None:
    document = Document()
    document.add_paragraph("一级条目")
    document.add_paragraph("二级条目")
    source = io.BytesIO()
    document.save(source)
    session = SimpleNamespace(
        document_v3={
            "styles": [
                {"id": "normal", "name": "正文", "kind": "paragraph", "visible": True, "properties": {}},
                {
                    "id": "outline-item",
                    "name": "大纲条目",
                    "kind": "paragraph",
                    "basedOn": "normal",
                    "visible": True,
                    "properties": {"numberingId": "thesis-outline", "numberingLevel": 1},
                },
            ],
            "numbering": [
                {
                    "id": "thesis-outline",
                    "name": "论文多级编号",
                    "levels": [
                        {"level": 0, "format": "decimal", "text": "%1", "startAt": 1, "leftIndentEm": 0, "hangingIndentEm": 0},
                        {"level": 1, "format": "decimal", "text": "%1.%2", "startAt": 1, "leftIndentEm": 2, "hangingIndentEm": 0},
                    ],
                }
            ],
            "sections": [
                {
                    "content": [
                        {"id": "p1", "type": "paragraph", "styleId": "outline-item", "content": [{"type": "text", "text": "一级条目"}]},
                        {"id": "p2", "type": "paragraph", "styleId": "outline-item", "attrs": {"numberingId": "thesis-outline", "numberingLevel": 1}, "content": [{"type": "text", "text": "二级条目"}]},
                    ]
                }
            ],
        }
    )

    result = ExportService._apply_document_v3_styles(source.getvalue(), session)
    exported = Document(io.BytesIO(result))
    numbering_xml = exported.part.numbering_part.element.xml

    assert 'w:multiLevelType w:val="multilevel"' in numbering_xml
    assert 'w:lvlText w:val="%1.%2"' in numbering_xml
    assert 'w:ilvl w:val="1"' in exported.styles["大纲条目"]._element.xml
    assert 'w:ilvl w:val="1"' in exported.paragraphs[1]._p.xml
