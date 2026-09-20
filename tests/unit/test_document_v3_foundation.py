from __future__ import annotations

import io

from fastapi.testclient import TestClient
from docx import Document

import writing_agent.web.app_v2 as app_v2
from writing_agent.v3 import DocumentCommand, DocumentV3, create_default_registry, migrate_doc_ir
from writing_agent.v3.document_commands import CommandTarget
from writing_agent.v3.document_model import BlockNode, to_plain_text, validate_unique_ids


def _legacy_doc() -> dict:
    return {
        "title": "测试论文",
        "sections": [
            {
                "id": "section-intro",
                "title": "绪论",
                "level": 1,
                "blocks": [{"id": "paragraph-1", "type": "paragraph", "text": "研究背景"}],
            }
        ],
    }


def test_legacy_doc_ir_migrates_with_stable_ids_and_named_styles() -> None:
    document = migrate_doc_ir(_legacy_doc())

    assert document.schema_version == 3
    assert document.sections[0].content[0].id == "section-intro"
    assert document.sections[0].content[0].style_id == "heading-1"
    assert document.sections[0].content[1].id == "paragraph-1"
    assert not validate_unique_ids(document)


def test_user_and_ai_share_the_same_validated_style_command() -> None:
    registry = create_default_registry()
    document = migrate_doc_ir(_legacy_doc())

    command = DocumentCommand(
        type="apply_style",
        target=CommandTarget(kind="nodes", node_ids=["paragraph-1"]),
        params={"style_id": "heading-2"},
        source="ai",
        review_mode="direct",
    )
    result = registry.execute(document, command)

    assert result.ok
    assert result.changed
    assert result.document is not None
    block = result.document.sections[0].content[1]
    assert block.type == "heading"
    assert block.style_id == "heading-2"
    assert block.attrs["level"] == 2
    assert document.sections[0].content[1].type == "paragraph"


def test_command_is_atomic_when_any_target_is_missing() -> None:
    registry = create_default_registry()
    document = migrate_doc_ir(_legacy_doc())
    command = DocumentCommand(
        type="apply_style",
        target=CommandTarget(kind="nodes", node_ids=["paragraph-1", "missing"]),
        params={"style_id": "heading-1"},
        source="user",
    )

    result = registry.execute(document, command)

    assert not result.ok
    assert not result.changed
    assert document.sections[0].content[1].type == "paragraph"


def test_registry_exposes_json_schema_for_ai_tools() -> None:
    schema = create_default_registry().tool_schema()

    assert schema["title"] == "DocumentCommand"
    assert "target" in schema["properties"]
    assert "params" in schema["properties"]


def test_word_style_properties_round_trip_through_the_canonical_model() -> None:
    document = migrate_doc_ir(_legacy_doc())
    document.styles[-1].properties.shading_color = "#f5f7fa"
    document.styles[-1].properties.letter_spacing_pt = 0.5
    document.styles[-1].properties.border_style = "solid"
    document.styles[-1].properties.tab_stops = [{"positionEm": 4, "alignment": "left"}]

    payload = document.model_dump(mode="json", by_alias=True)
    restored = DocumentV3.model_validate(payload)

    properties = restored.styles[-1].properties
    assert properties.shading_color == "#f5f7fa"
    assert properties.letter_spacing_pt == 0.5
    assert properties.border_style == "solid"
    assert properties.tab_stops == [{"positionEm": 4, "alignment": "left"}]


def test_plain_text_bridge_preserves_structured_tables_for_export() -> None:
    document = migrate_doc_ir(_legacy_doc())
    document.sections[0].content.append(
        BlockNode(
            id="table-1",
            type="table",
            attrs={
                "table": {
                    "caption": "测试表",
                    "columns": ["项目", "结果"],
                    "rows": [["保存", "成功"]],
                }
            },
        )
    )

    text = to_plain_text(document)

    assert '[[TABLE:{"caption":"测试表","columns":["项目","结果"],"rows":[["保存","成功"]]}]]' in text


def test_saved_document_v3_table_is_present_in_docx_export() -> None:
    session = app_v2.store.create()
    document = migrate_doc_ir(_legacy_doc())
    document.sections[0].content.append(
        BlockNode(
            id="table-export",
            type="table",
            attrs={"table": {"columns": ["项目", "结果"], "rows": [["导出", "成功"]]}},
        )
    )
    client = TestClient(app_v2.app)
    try:
        save = client.post(
            f"/api/doc/{session.id}/save",
            json={"document_v3": document.model_dump(mode="json", by_alias=True)},
        )
        assert save.status_code == 200

        response = client.get(f"/download/{session.id}.docx")
        assert response.status_code == 200
        exported = Document(io.BytesIO(response.content))
        assert len(exported.tables) == 1
        cells = [cell.text for row in exported.tables[0].rows for cell in row.cells]
        assert cells == ["项目", "结果", "导出", "成功"]
    finally:
        app_v2.store.delete(session.id)


def test_document_v3_command_api_persists_an_atomic_ai_edit() -> None:
    session = app_v2.store.create()
    session.doc_ir = _legacy_doc()
    app_v2.store.put(session)
    client = TestClient(app_v2.app)
    try:
        schema_response = client.get("/api/document-v3/tool-schema")
        assert schema_response.status_code == 200
        assert "apply_style" in schema_response.json()["command_types"]

        response = client.post(
            f"/api/doc/{session.id}/document-v3/commands",
            json={
                "type": "apply_style",
                "target": {"kind": "nodes", "node_ids": ["paragraph-1"]},
                "params": {"style_id": "heading-2"},
                "source": "ai",
                "review_mode": "direct",
            },
        )
        assert response.status_code == 200
        stored = app_v2.store.get(session.id)
        assert stored is not None
        assert stored.document_v3["schemaVersion"] == 3
        block = stored.document_v3["sections"][0]["content"][1]
        assert block["styleId"] == "heading-2"
        assert stored.doc_text == "# 绪论\n\n## 研究背景"
    finally:
        app_v2.store.delete(session.id)
