from __future__ import annotations

import base64
import io

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from fastapi.testclient import TestClient

import writing_agent.web.app_v2 as app_v2
from writing_agent.v3.docx_import import import_docx_bytes
from writing_agent.v3.document_model import iter_blocks, to_plain_text


PNG_1X1 = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAusB9WlL4SAAAAAASUVORK5CYII="
)


def _structured_docx() -> bytes:
    document = Document()
    document.core_properties.title = "结构化导入测试"
    custom = document.styles.add_style("论文正文", WD_STYLE_TYPE.PARAGRAPH)
    custom.base_style = document.styles["Normal"]
    custom.font.name = "Microsoft YaHei"
    custom.font.size = Pt(11)
    custom.paragraph_format.line_spacing = 1.5
    custom.paragraph_format.first_line_indent = Pt(22)

    document.add_heading("第一章 绪论", level=1)
    body = document.add_paragraph(style=custom)
    body.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    body.add_run("正文加粗").bold = True
    body.add_run("与普通文字")
    picture = io.BytesIO(PNG_1X1)
    body.add_run().add_picture(picture, width=Inches(0.2))
    document.add_paragraph("第一项", style="List Number")
    document.add_paragraph("第二项", style="List Number")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "项目"
    table.cell(0, 1).text = "结果"
    table.cell(1, 0).text = "导入"
    table.cell(1, 1).text = "成功"
    document.sections[0].header.paragraphs[0].text = "论文页眉"
    document.sections[0].footer.paragraphs[0].text = "论文页脚"
    document.add_section(WD_SECTION.NEW_PAGE)
    document.add_heading("第二章 设计", level=1)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def test_docx_import_preserves_editable_structure_and_styles() -> None:
    document = import_docx_bytes(_structured_docx(), title="测试论文")
    blocks = list(iter_blocks(document))
    types = [block.type for block in blocks]

    assert document.title == "测试论文"
    assert len(document.sections) == 2
    assert {"heading", "paragraph", "orderedList", "listItem", "table", "figure"} <= set(types)
    custom = next(style for style in document.styles if style.name == "论文正文")
    assert custom.properties.font_family == "Microsoft YaHei"
    assert custom.properties.font_size_pt == 11
    assert custom.properties.line_spacing == 1.5
    assert document.sections[0].header_footer.header[0].content[0].text == "论文页眉"
    assert document.sections[0].header_footer.footer[0].content[0].text == "论文页脚"
    assert any(item.get("id", "").startswith("word-num-") for item in document.numbering)
    assert not any("data:image" in line for line in to_plain_text(document, include_resource_data=False).splitlines())


def test_docx_import_endpoint_saves_document_v3_without_duplicating_image_data_in_text() -> None:
    session = app_v2.store.create()
    client = TestClient(app_v2.app)
    original = _structured_docx()
    try:
        response = client.post(
            f"/api/doc/{session.id}/import",
            files={"file": ("structured.docx", original, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        )
        assert response.status_code == 200, response.text
        payload = response.json()
        assert payload["document_v3"]["schemaVersion"] == 3
        assert payload["imported"]["tables"] == 1
        stored = app_v2.store.get(session.id)
        assert stored is not None
        assert stored.document_v3["sections"]
        assert "data:image" not in stored.doc_text
        assert len(stored.doc_text) < 10_000
        assert stored.import_source_fingerprint
        source_path = app_v2.store.source_document_path(session.id)
        assert source_path is not None and source_path.read_bytes() == original

        export = client.get(f"/download/{session.id}.docx")
        assert export.status_code == 200, export.text
        assert export.content == source_path.read_bytes()
        assert export.headers["x-docx-export-backend"] == "lossless_import_roundtrip"
    finally:
        app_v2.store.delete(session.id)
        assert app_v2.store.source_document_path(session.id) is None


def test_edited_import_patches_original_docx_instead_of_flattening_it() -> None:
    session = app_v2.store.create()
    client = TestClient(app_v2.app)
    try:
        response = client.post(
            f"/api/doc/{session.id}/import",
            files={"file": ("structured.docx", _structured_docx(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        )
        payload = response.json()["document_v3"]
        heading = next(
            block
            for section in payload["sections"]
            for block in section["content"]
            if block["type"] == "heading"
        )
        heading["content"] = [{"type": "text", "text": "第一章 修改后的绪论", "marks": [], "attrs": {}}]
        body = next(
            block
            for section in payload["sections"]
            for block in section["content"]
            if block["type"] == "paragraph" and block.get("styleId", "").startswith("word-style-")
        )
        body["content"] = [{"type": "text", "text": "修改后的正文，图片仍应保留。", "marks": [], "attrs": {}}]
        save = client.post(f"/api/doc/{session.id}/save", json={"document_v3": payload})
        assert save.status_code == 200, save.text

        export = client.get(f"/download/{session.id}.docx")
        assert export.status_code == 200, export.text
        assert export.headers["x-docx-export-backend"] == "structured_import_patch"
        result = Document(io.BytesIO(export.content))
        assert result.paragraphs[0].text == "第一章 修改后的绪论"
        assert any(paragraph.text == "修改后的正文，图片仍应保留。" for paragraph in result.paragraphs)
        assert len(result.tables) == 1
        assert len(result.inline_shapes) == 1
        assert len(result.sections) == 2
        assert result.sections[0].header.paragraphs[0].text == "论文页眉"
    finally:
        app_v2.store.delete(session.id)

