"""Structure-preserving DOCX importer for the canonical Document V3 model."""

from __future__ import annotations

import base64
import hashlib
import io
import json
import re
from copy import deepcopy
from typing import Any
from uuid import uuid4

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from .document_model import (
    BlockNode,
    DocumentV3,
    HeaderFooterDefinition,
    InlineNode,
    PageLayout,
    PageNumberDefinition,
    SectionV3,
    StyleDefinition,
    StyleProperties,
    default_numbering,
    default_styles,
)


def _id(prefix: str) -> str:
    return f"{prefix}_{uuid4().hex}"


def _style_key(style) -> str:
    builtin = {
        "Normal": "normal",
        "Title": "title",
        "Subtitle": "subtitle",
        "Quote": "quote",
        "Caption": "caption",
        "Code": "code",
        "论文题目": "title",
        "章标题": "heading-1",
        "节标题": "heading-2",
        "表格头": "caption",
    }
    style_id = str(getattr(style, "style_id", "") or getattr(style, "name", "") or "normal")
    style_name = str(getattr(style, "name", "") or "")
    if style_id in builtin:
        return builtin[style_id]
    if style_name in builtin:
        return builtin[style_name]
    heading = re.fullmatch(r"Heading\s*([1-6])", style_id, flags=re.IGNORECASE)
    if heading:
        return f"heading-{heading.group(1)}"
    safe = re.sub(r"[^a-z0-9]+", "-", style_id.lower()).strip("-")
    return f"word-style-{safe or uuid4().hex[:8]}"


def _color(value) -> str | None:
    rgb = getattr(value, "rgb", None)
    return f"#{rgb}" if rgb is not None else None


def _pt(value) -> float | None:
    return round(float(value.pt), 3) if value is not None else None


def _alignment(value) -> str | None:
    return {
        WD_ALIGN_PARAGRAPH.LEFT: "left",
        WD_ALIGN_PARAGRAPH.CENTER: "center",
        WD_ALIGN_PARAGRAPH.RIGHT: "right",
        WD_ALIGN_PARAGRAPH.JUSTIFY: "justify",
        WD_ALIGN_PARAGRAPH.DISTRIBUTE: "justify",
    }.get(value)


def _style_properties(style) -> StyleProperties:
    font = style.font
    fmt = style.paragraph_format
    font_size = _pt(font.size) or 12.0
    line_spacing = fmt.line_spacing
    if hasattr(line_spacing, "pt"):
        line_spacing = round(float(line_spacing.pt) / font_size, 3)
    elif line_spacing is not None:
        line_spacing = round(float(line_spacing), 3)
    p_pr = style._element.pPr
    outline_level = None
    shading_color = None
    border_style = None
    border_color = None
    border_width = None
    if p_pr is not None:
        outline = p_pr.find(qn("w:outlineLvl"))
        if outline is not None:
            outline_level = int(outline.get(qn("w:val"), "0")) + 1
        shading = p_pr.find(qn("w:shd"))
        if shading is not None and shading.get(qn("w:fill"), "auto") not in {"auto", "nil"}:
            shading_color = f"#{shading.get(qn('w:fill'))}"
        borders = p_pr.find(qn("w:pBdr"))
        edge = borders.find(qn("w:bottom")) if borders is not None else None
        if edge is not None:
            border_style = {"single": "solid", "dashed": "dashed", "double": "double", "nil": "none"}.get(edge.get(qn("w:val"), ""), "solid")
            edge_color = edge.get(qn("w:color"), "")
            border_color = f"#{edge_color}" if re.fullmatch(r"[0-9A-Fa-f]{6}", edge_color) else None
            border_width = round(float(edge.get(qn("w:sz"), "4")) / 8, 3)
    tab_stops = []
    for tab in fmt.tab_stops:
        alignment = {0: "left", 1: "center", 2: "right", 3: "decimal"}.get(int(tab.alignment), "left")
        tab_stops.append({"positionEm": round(float(tab.position.pt) / font_size, 3), "alignment": alignment})
    return StyleProperties(
        font_family=font.name,
        font_size_pt=_pt(font.size),
        bold=font.bold,
        italic=font.italic,
        underline=bool(font.underline) if font.underline is not None else None,
        color=_color(font.color),
        alignment=_alignment(fmt.alignment),
        line_spacing=line_spacing,
        first_line_indent_em=round(float(fmt.first_line_indent.pt) / font_size, 3) if fmt.first_line_indent is not None else None,
        left_indent_em=round(float(fmt.left_indent.pt) / font_size, 3) if fmt.left_indent is not None else None,
        right_indent_em=round(float(fmt.right_indent.pt) / font_size, 3) if fmt.right_indent is not None else None,
        space_before_pt=_pt(fmt.space_before),
        space_after_pt=_pt(fmt.space_after),
        outline_level=outline_level,
        keep_with_next=fmt.keep_with_next,
        keep_lines_together=fmt.keep_together,
        page_break_before=fmt.page_break_before,
        border_color=border_color,
        border_width_pt=border_width,
        border_style=border_style,
        shading_color=shading_color,
        tab_stops=tab_stops or None,
    )


def _import_styles(document) -> tuple[list[StyleDefinition], dict[str, str]]:
    result = deepcopy(default_styles())
    by_id = {style.id: index for index, style in enumerate(result)}
    mapping: dict[str, str] = {}
    paragraph_styles = [style for style in document.styles if style.type == WD_STYLE_TYPE.PARAGRAPH]
    for style in paragraph_styles:
        mapping[str(style.style_id)] = _style_key(style)
        mapping[str(style.name)] = _style_key(style)
    for style in paragraph_styles:
        style_id = _style_key(style)
        definition = StyleDefinition(
            id=style_id,
            name=str(style.name or style.style_id or style_id),
            kind="paragraph",
            based_on=_style_key(style.base_style) if style.base_style is not None else None,
            next_style=_style_key(style.next_paragraph_style) if style.next_paragraph_style is not None else style_id,
            visible=not bool(getattr(style, "hidden", False)),
            properties=_style_properties(style),
        )
        if style_id in by_id:
            result[by_id[style_id]] = definition
        else:
            by_id[style_id] = len(result)
            result.append(definition)
    return result, mapping


def _parse_numbering(document) -> tuple[list[dict[str, Any]], dict[str, dict[str, Any]]]:
    try:
        root = document.part.numbering_part.element
    except Exception:
        defaults = default_numbering()
        return defaults, {}
    abstracts = {item.get(qn("w:abstractNumId")): item for item in root.findall(qn("w:abstractNum"))}
    definitions: list[dict[str, Any]] = []
    lookup: dict[str, dict[str, Any]] = {}
    for number in root.findall(qn("w:num")):
        num_id = str(number.get(qn("w:numId")) or "")
        abstract_ref = number.find(qn("w:abstractNumId"))
        abstract = abstracts.get(abstract_ref.get(qn("w:val")) if abstract_ref is not None else "")
        if not num_id or abstract is None:
            continue
        levels: list[dict[str, Any]] = []
        for raw_level in abstract.findall(qn("w:lvl")):
            level = max(0, min(8, int(raw_level.get(qn("w:ilvl"), "0"))))
            fmt_node = raw_level.find(qn("w:numFmt"))
            text_node = raw_level.find(qn("w:lvlText"))
            start_node = raw_level.find(qn("w:start"))
            fmt = fmt_node.get(qn("w:val"), "decimal") if fmt_node is not None else "decimal"
            fmt = fmt if fmt in {"bullet", "decimal", "lowerLetter", "upperLetter", "lowerRoman", "upperRoman"} else "decimal"
            text = text_node.get(qn("w:val"), f"%{level + 1}.") if text_node is not None else f"%{level + 1}."
            indent = raw_level.find(f"{qn('w:pPr')}/{qn('w:ind')}")
            left = int(indent.get(qn("w:left"), "0")) / 240 if indent is not None else 2 + level * 2
            hanging = int(indent.get(qn("w:hanging"), "0")) / 240 if indent is not None else 1
            item = {
                "level": level,
                "format": fmt,
                "text": text,
                "startAt": max(1, int(start_node.get(qn("w:val"), "1"))) if start_node is not None else 1,
                "leftIndentEm": round(left, 3),
                "hangingIndentEm": round(hanging, 3),
            }
            if fmt == "bullet":
                item["bulletChar"] = text
            levels.append(item)
        if not levels:
            continue
        definition_id = f"word-num-{num_id}"
        definition = {"id": definition_id, "name": f"Word 编号 {num_id}", "levels": sorted(levels, key=lambda item: item["level"])}
        definitions.append(definition)
        lookup[num_id] = definition
    return definitions or default_numbering(), lookup


def _paragraph_numbering(paragraph: Paragraph, lookup: dict[str, dict[str, Any]]) -> tuple[str, int, str] | None:
    num_pr = paragraph._p.pPr.numPr if paragraph._p.pPr is not None else None
    if num_pr is None and paragraph.style is not None and paragraph.style._element.pPr is not None:
        num_pr = paragraph.style._element.pPr.find(qn("w:numPr"))
    if num_pr is None:
        return None
    num_id_node = num_pr.find(qn("w:numId"))
    level_node = num_pr.find(qn("w:ilvl"))
    num_id = num_id_node.get(qn("w:val"), "") if num_id_node is not None else ""
    level = max(0, min(8, int(level_node.get(qn("w:val"), "0")))) if level_node is not None else 0
    definition = lookup.get(num_id)
    if not definition:
        return None
    level_definition = next((item for item in definition["levels"] if item["level"] == level), definition["levels"][0])
    return str(definition["id"]), level, str(level_definition.get("format") or "decimal")


def _inline_content(paragraph: Paragraph) -> list[InlineNode]:
    content: list[InlineNode] = []
    for run in paragraph.runs:
        marks: list[dict[str, Any]] = []
        if run.bold:
            marks.append({"type": "bold"})
        if run.italic:
            marks.append({"type": "italic"})
        if run.underline:
            marks.append({"type": "underline"})
        if run.font.strike:
            marks.append({"type": "strike"})
        if run.font.subscript:
            marks.append({"type": "subscript"})
        if run.font.superscript:
            marks.append({"type": "superscript"})
        text_attrs: dict[str, Any] = {}
        if run.font.name:
            text_attrs["fontFamily"] = run.font.name
        if run.font.size is not None:
            text_attrs["fontSize"] = f"{round(float(run.font.size.pt), 3)}pt"
        run_color = _color(run.font.color)
        if run_color:
            text_attrs["color"] = run_color
        if text_attrs:
            marks.append({"type": "textStyle", "attrs": text_attrs})
        pieces = str(run.text or "").split("\n")
        for index, piece in enumerate(pieces):
            if piece:
                content.append(InlineNode(type="text", text=piece, marks=deepcopy(marks)))
            if index < len(pieces) - 1:
                content.append(InlineNode(type="hardBreak"))
    return content


def _paragraph_attrs(paragraph: Paragraph) -> dict[str, Any]:
    fmt = paragraph.paragraph_format
    font_size = _pt(paragraph.style.font.size if paragraph.style is not None else None) or 12.0
    attrs: dict[str, Any] = {}
    alignment = _alignment(fmt.alignment)
    if alignment:
        attrs["textAlign"] = alignment
    line_spacing = fmt.line_spacing
    if hasattr(line_spacing, "pt"):
        attrs["lineSpacing"] = round(float(line_spacing.pt) / font_size, 3)
    elif line_spacing is not None:
        attrs["lineSpacing"] = round(float(line_spacing), 3)
    for source, target in (
        (fmt.first_line_indent, "firstLineIndentEm"),
        (fmt.left_indent, "leftIndentEm"),
        (fmt.right_indent, "rightIndentEm"),
    ):
        if source is not None:
            attrs[target] = round(float(source.pt) / font_size, 3)
    if fmt.space_before is not None:
        attrs["spaceBeforePt"] = _pt(fmt.space_before)
    if fmt.space_after is not None:
        attrs["spaceAfterPt"] = _pt(fmt.space_after)
    if fmt.keep_with_next is not None:
        attrs["keepWithNext"] = bool(fmt.keep_with_next)
    if fmt.keep_together is not None:
        attrs["keepLinesTogether"] = bool(fmt.keep_together)
    if fmt.page_break_before is not None:
        attrs["pageBreakBefore"] = bool(fmt.page_break_before)
    return attrs


def _paragraph_block(paragraph: Paragraph, style_mapping: dict[str, str], source_index: int | None = None) -> BlockNode:
    style_id = style_mapping.get(str(paragraph.style.style_id), _style_key(paragraph.style)) if paragraph.style is not None else "normal"
    heading = re.fullmatch(r"heading-([1-6])", style_id)
    if paragraph.style is not None and str(paragraph.style.name or "") == "节标题":
        numbered = re.match(r"^\s*(\d+(?:\.\d+){0,5})(?:\s|[、.．])", paragraph.text or "")
        if numbered:
            inferred_level = min(6, numbered.group(1).count(".") + 1)
            style_id = f"heading-{inferred_level}"
            heading = re.fullmatch(r"heading-([1-6])", style_id)
    block_type = "heading" if heading else "paragraph"
    attrs = _paragraph_attrs(paragraph)
    if source_index is not None:
        attrs["sourceParagraphIndex"] = source_index
    if heading:
        attrs["level"] = int(heading.group(1))
    return BlockNode(id=_id(block_type), type=block_type, style_id=style_id, attrs=attrs, content=_inline_content(paragraph))


def _paragraph_figures(paragraph: Paragraph, source_index: int | None = None) -> list[BlockNode]:
    figures: list[BlockNode] = []
    for blip in paragraph._p.xpath(".//a:blip"):
        relationship_id = blip.get(qn("r:embed"))
        part = paragraph.part.related_parts.get(relationship_id) if relationship_id else None
        if part is None or not getattr(part, "blob", b""):
            continue
        content_type = str(getattr(part, "content_type", "application/octet-stream"))
        source = f"data:{content_type};base64,{base64.b64encode(part.blob).decode('ascii')}"
        figures.append(
            BlockNode(
                id=_id("figure"),
                type="figure",
                attrs={
                    "sourceParagraphIndex": source_index,
                    "figure": {"src": source, "alt": "从 Word 导入的图片", "caption": "", "widthPercent": 100, "alignment": "center", "wrap": "inline"},
                },
            )
        )
    return figures


def _table_block(table: Table, source_index: int | None = None) -> BlockNode:
    rows: list[list[str]] = []
    cells: list[list[dict[str, Any]]] = []
    for row_index, row in enumerate(table.rows):
        row_text: list[str] = []
        row_cells: list[dict[str, Any]] = []
        for cell in row.cells:
            text = "\n".join(paragraph.text for paragraph in cell.paragraphs).strip()
            grid_span = cell._tc.tcPr.gridSpan
            colspan = int(grid_span.val) if grid_span is not None else 1
            vertical = cell._tc.tcPr.vMerge
            row_text.append(text)
            row_cells.append(
                {
                    "text": text,
                    "type": "header" if row_index == 0 else "cell",
                    "colspan": colspan,
                    "rowspan": 1,
                    "verticalMerge": str(vertical.val or "continue") if vertical is not None else None,
                }
            )
        rows.append(row_text)
        cells.append(row_cells)
    repeat_header = bool(table.rows and table.rows[0]._tr.xpath("./w:trPr/w:tblHeader"))
    alignment = {0: "left", 1: "center", 2: "right"}.get(int(table.alignment) if table.alignment is not None else 0, "left")
    return BlockNode(
        id=_id("table"),
        type="table",
        attrs={
            "sourceTableIndex": source_index,
            "table": {
                "caption": "",
                "columns": rows[0] if rows else [],
                "rows": rows[1:] if rows else [],
                "cells": cells,
                "repeatHeader": repeat_header,
                "alignment": alignment,
                "widthPercent": 100,
            }
        },
    )


def _header_footer(section, style_mapping: dict[str, str]) -> HeaderFooterDefinition:
    def blocks(container) -> list[BlockNode]:
        return [_paragraph_block(paragraph, style_mapping) for paragraph in container.paragraphs if paragraph.text.strip()]

    page_number = PageNumberDefinition(enabled=True, format="arabic", position="footer", alignment="center")
    sect_pr = section._sectPr
    page_number_type = sect_pr.find(qn("w:pgNumType"))
    if page_number_type is not None:
        page_number.start_at = int(page_number_type.get(qn("w:start"), "1"))
        value = page_number_type.get(qn("w:fmt"), "decimal")
        page_number.format = {"decimal": "arabic", "lowerRoman": "lowerRoman", "upperRoman": "upperRoman", "lowerLetter": "lowerLetter", "upperLetter": "upperLetter"}.get(value, "arabic")
    return HeaderFooterDefinition(
        link_header_to_previous=bool(section.header.is_linked_to_previous),
        link_footer_to_previous=bool(section.footer.is_linked_to_previous),
        different_first_page=bool(section.different_first_page_header_footer),
        different_odd_even=bool(document_settings_even_odd(section)),
        header=blocks(section.header),
        footer=blocks(section.footer),
        first_page_header=blocks(section.first_page_header),
        first_page_footer=blocks(section.first_page_footer),
        even_page_header=blocks(section.even_page_header),
        even_page_footer=blocks(section.even_page_footer),
        page_number=page_number,
    )


def document_settings_even_odd(section) -> bool:
    settings = section._document_part.document.settings.element
    return settings.find(qn("w:evenAndOddHeaders")) is not None


def _layout(section) -> PageLayout:
    width = round(float(section.page_width.mm), 3)
    height = round(float(section.page_height.mm), 3)
    landscape = width > height
    known = {(210.0, 297.0): "A4", (148.0, 210.0): "A5", (297.0, 420.0): "A3", (215.9, 279.4): "Letter"}
    portrait = (min(width, height), max(width, height))
    page_size = next((name for (w, h), name in known.items() if abs(portrait[0] - w) < 1 and abs(portrait[1] - h) < 1), "custom")
    return PageLayout(
        page_size=page_size,
        width_mm=width if page_size == "custom" else None,
        height_mm=height if page_size == "custom" else None,
        orientation="landscape" if landscape else "portrait",
        margin_top_mm=round(float(section.top_margin.mm), 3),
        margin_right_mm=round(float(section.right_margin.mm), 3),
        margin_bottom_mm=round(float(section.bottom_margin.mm), 3),
        margin_left_mm=round(float(section.left_margin.mm), 3),
        columns=1,
    )


def _append_list_paragraph(
    section_content: list[BlockNode],
    paragraph: Paragraph,
    block: BlockNode,
    numbering: tuple[str, int, str],
    list_state: dict[str, Any],
) -> None:
    numbering_id, level, number_format = numbering
    list_type = "bulletList" if number_format == "bullet" else "orderedList"
    if list_state.get("numbering_id") != numbering_id or list_state.get("list_type") != list_type or level == 0 and list_state.get("last_level", 0) > 0:
        list_state.clear()
    if not list_state:
        root = BlockNode(id=_id(list_type), type=list_type, attrs={"numberingId": numbering_id, "numberingLevel": 0}, content=[])
        section_content.append(root)
        list_state.update({"numbering_id": numbering_id, "list_type": list_type, "root": root, "items": {}, "lists": {0: root}, "last_level": 0})
    if level not in list_state["lists"]:
        parent_level = max((candidate for candidate in list_state["items"] if candidate < level), default=0)
        parent_item = list_state["items"].get(parent_level)
        if parent_item is None:
            level = 0
        else:
            nested = BlockNode(id=_id(list_type), type=list_type, attrs={"numberingId": numbering_id, "numberingLevel": level}, content=[])
            parent_item.content.append(nested)
            list_state["lists"][level] = nested
    target_list = list_state["lists"].get(level, list_state["root"])
    block.attrs["numberingId"] = numbering_id
    block.attrs["numberingLevel"] = level
    item = BlockNode(id=_id("listItem"), type="listItem", attrs={"numberingId": numbering_id, "numberingLevel": level}, content=[block])
    target_list.content.append(item)
    list_state["items"][level] = item
    for candidate in list(list_state["items"]):
        if candidate > level:
            list_state["items"].pop(candidate, None)
            list_state["lists"].pop(candidate, None)
    list_state["last_level"] = level


def import_docx_bytes(payload: bytes, *, title: str = "") -> DocumentV3:
    """Import a DOCX without flattening its editable document structure."""
    document = Document(io.BytesIO(payload))
    styles, style_mapping = _import_styles(document)
    numbering, numbering_lookup = _parse_numbering(document)
    sections: list[SectionV3] = []
    word_sections = list(document.sections)

    def new_section(index: int) -> SectionV3:
        source = word_sections[min(index, len(word_sections) - 1)]
        return SectionV3(
            id=_id("section"),
            break_type="nextPage",
            layout=_layout(source),
            header_footer=_header_footer(source, style_mapping),
            content=[],
        )

    sections.append(new_section(0))
    section_index = 0
    list_state: dict[str, Any] = {}
    paragraph_indices = {paragraph._p: index for index, paragraph in enumerate(document.paragraphs)}
    table_indices = {table._tbl: index for index, table in enumerate(document.tables)}
    represented_paragraphs: list[int] = []
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            paragraph = Paragraph(child, document)
            source_index = paragraph_indices.get(child)
            paragraph_style_id = str(paragraph.style.style_id or "") if paragraph.style is not None else ""
            paragraph_style_name = str(paragraph.style.name or "") if paragraph.style is not None else ""
            if re.fullmatch(r"TOC[1-9]", paragraph_style_id, flags=re.IGNORECASE) or re.fullmatch(r"toc\s*[1-9]", paragraph_style_name, flags=re.IGNORECASE):
                continue
            instruction = " ".join(paragraph._p.xpath(".//w:instrText/text()"))
            if re.search(r"(?:^|\s)TOC(?:\s|$)", instruction, flags=re.IGNORECASE):
                list_state.clear()
                sections[-1].content.append(BlockNode(id=_id("toc"), type="tableOfContents", attrs={"levels": [1, 2, 3], "sourceParagraphIndex": source_index}))
                if source_index is not None:
                    represented_paragraphs.append(source_index)
            else:
                block = _paragraph_block(paragraph, style_mapping, source_index)
                paragraph_numbering = _paragraph_numbering(paragraph, numbering_lookup)
                has_page_break = bool(paragraph._p.xpath(".//w:br[@w:type='page']"))
                if paragraph_numbering and (paragraph.text.strip() or block.content):
                    _append_list_paragraph(sections[-1].content, paragraph, block, paragraph_numbering, list_state)
                    if source_index is not None:
                        represented_paragraphs.append(source_index)
                else:
                    list_state.clear()
                    if paragraph.text.strip() or block.content:
                        sections[-1].content.append(block)
                        if source_index is not None:
                            represented_paragraphs.append(source_index)
                    figures = _paragraph_figures(paragraph, source_index)
                    sections[-1].content.extend(figures)
                    if figures and source_index is not None and source_index not in represented_paragraphs:
                        represented_paragraphs.append(source_index)
                if has_page_break:
                    sections[-1].content.append(BlockNode(id=_id("pageBreak"), type="pageBreak"))
            if paragraph._p.pPr is not None and paragraph._p.pPr.sectPr is not None and section_index + 1 < len(word_sections):
                section_index += 1
                sections.append(new_section(section_index))
                list_state.clear()
        elif child.tag == qn("w:tbl"):
            list_state.clear()
            sections[-1].content.append(_table_block(Table(child, document), table_indices.get(child)))
    if not any(section.content for section in sections):
        sections[0].content.append(BlockNode(id=_id("paragraph"), type="paragraph", style_id="normal"))
    name = title.strip() or str(document.core_properties.title or "").strip() or "导入的 Word 文档"
    return DocumentV3(
        title=name,
        metadata={
            "importedFrom": "docx",
            "sourceParagraphs": len(document.paragraphs),
            "sourceTables": len(document.tables),
            "sourceSections": len(document.sections),
            "representedParagraphs": sorted(set(represented_paragraphs)),
        },
        styles=styles,
        numbering=numbering,
        sections=sections,
    )


def document_v3_fingerprint(document: DocumentV3 | dict[str, Any]) -> str:
    """Return a stable digest of editable V3 content.

    Runtime metadata and the generated document id do not affect whether an
    imported DOCX can be returned byte-for-byte.
    """
    parsed = document if isinstance(document, DocumentV3) else DocumentV3.model_validate(document)
    payload = parsed.model_dump(mode="json", by_alias=True)
    payload.pop("id", None)
    payload.pop("metadata", None)
    encoded = json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()

